import os
import io
import re
import html
import requests
import streamlit as st
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from fpdf import FPDF

API_URL = os.getenv("API_URL", "http://localhost:8000/generate")
LOGO = os.path.join(os.path.dirname(__file__), "..", "Image", "Logo.png")


def sanitize_text(text):
    rep = {"\u2018": "'", "\u2019": "'", "\u201c": '"', "\u201d": '"',
           "\u2013": "-", "\u2014": "-", "\u2022": "-", "\u2026": "..."}
    for k, v in rep.items():
        text = text.replace(k, v)
    text = text.replace("**", "")
    text = re.sub(r"^#+\s*", "", text, flags=re.M)
    return text.encode("latin-1", "ignore").decode("latin-1")


def format_html_preview(text):
    return html.escape(text).replace("\n", "<br>")


def format_docx(text, doc_type, terms):
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)
    if os.path.exists(LOGO):
        doc.add_picture(LOGO, width=Inches(1.5))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_heading(doc_type, 0).alignment = WD_ALIGN_PARAGRAPH.CENTER
    for line in text.split("\n"):
        if line.strip():
            doc.add_paragraph(line.strip())
    items = [t.strip() for t in terms.split(";") if t.strip()]
    if items:
        doc.add_heading("Key Terms", 1)
        table = doc.add_table(rows=1, cols=2)
        table.style = "Table Grid"
        table.rows[0].cells[0].text = "No."
        table.rows[0].cells[1].text = "Term"
        for i, t in enumerate(items, 1):
            row = table.add_row().cells
            row[0].text = str(i)
            row[1].text = t
    footer = doc.sections[0].footer.paragraphs[0]
    footer.text = "LegalEase | All Rights Reserved"
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


class PDF(FPDF):
    def header(self):
        if os.path.exists(LOGO):
            self.image(LOGO, x=90, y=8, w=30)
            self.ln(25)

    def footer(self):
        self.set_y(-15)
        self.set_font("Helvetica", "I", 8)
        self.cell(0, 10, "LegalEase | All Rights Reserved", align="C")


def format_pdf(text, doc_type):
    pdf = PDF()
    pdf.add_page()
    pdf.set_font("Helvetica", "B", 16)
    pdf.multi_cell(0, 10, doc_type, align="C", new_x="LMARGIN", new_y="NEXT")
    pdf.ln(4)
    pdf.set_font("Helvetica", "", 11)
    for line in text.split("\n"):
        if line.strip():
            pdf.multi_cell(0, 6, line.strip(), new_x="LMARGIN", new_y="NEXT")
    return bytes(pdf.output())


# ---------------- UI ----------------
st.set_page_config(page_title="LegalEase", layout="centered")

if os.path.exists(LOGO):
    c1, c2, c3 = st.columns([1, 2, 1])
    with c2:
        st.image(LOGO, use_container_width=True)
st.markdown("<h2 style='text-align:center;'>AI Legal Document Generator</h2>",
            unsafe_allow_html=True)

document_type = st.text_input("Document Type (Ex: Agreement, Contract, NDA)")
parties = st.text_area("Parties Involved")
terms = st.text_area("Terms & Conditions (Use semicolons for bullet points)")
dates = st.text_input("Effective Date")

if st.button("Generate Document"):
    if not (document_type and parties and terms and dates):
        st.warning("Please fill all fields.")
    else:
        with st.spinner("Generating..."):
            try:
                r = requests.post(API_URL, json={
                    "document_type": document_type, "parties": parties,
                    "terms": terms, "dates": dates}, timeout=120)
                r.raise_for_status()
                st.session_state.doc = sanitize_text(r.json()["document"])
                st.session_state.doc_type = document_type
                st.session_state.terms = terms
                st.session_state.show_edit = False
            except Exception as e:
                st.error(f"Error: {e}")

if "doc" in st.session_state:
    st.success("Document Generated Successfully!")
    st.markdown(
        "<div style='background:#0f172a;padding:16px;border-radius:8px;"
        "max-height:350px;overflow-y:auto;color:#e2e8f0;'>"
        f"{format_html_preview(st.session_state.doc)}</div>",
        unsafe_allow_html=True)

    if st.button("Click to Edit Document"):
        st.session_state.show_edit = True
    if st.session_state.get("show_edit"):
        st.session_state.doc = st.text_area(
            "Edit Document Below:", st.session_state.doc, height=300)

    name = st.session_state.doc_type.replace(" ", "_").lower()
    st.download_button("Download as .TXT", st.session_state.doc,
                       file_name=f"{name}.txt")
    st.download_button("Download as .DOCX",
                       format_docx(st.session_state.doc, st.session_state.doc_type,
                                   st.session_state.terms),
                       file_name=f"{name}.docx")
    st.download_button("Download as .PDF",
                       format_pdf(st.session_state.doc, st.session_state.doc_type),
                       file_name=f"{name}.pdf")